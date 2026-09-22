from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SHOTS = ROOT / "screenshots"
OUT = ROOT / "output"
OUT.mkdir(parents=True, exist_ok=True)
DOCX_PATH = OUT / "Panduan-Admin-Mahaga-Widya-Cita.docx"

NAVY = "17324D"
BLUE = "2E74B5"
LIGHT_BLUE = "E8F1F8"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "667085"
GOLD = "B58A2A"
RED = "9B1C1C"
GREEN = "1F6B45"
WHITE = "FFFFFF"
BLACK = "111827"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, text, end):
        run._r.append(element)


def font(run, size=None, bold=None, color=None, italic=None, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.35)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(BLACK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15

for name, size, color, before, after in (
    ("Title", 28, NAVY, 0, 8),
    ("Subtitle", 13, MID_GRAY, 0, 16),
    ("Heading 1", 18, NAVY, 16, 8),
    ("Heading 2", 14, BLUE, 12, 6),
    ("Heading 3", 11.5, NAVY, 9, 4),
):
    style = styles[name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = name != "Subtitle"
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

for name in ("List Bullet", "List Number"):
    style = styles[name]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    style.paragraph_format.left_indent = Inches(0.38)
    style.paragraph_format.first_line_indent = Inches(-0.19)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.15

header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font(hp.add_run("MAHAGA WIDYA CITA  |  ADMIN CMS"), 8, True, MID_GRAY)
footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font(fp.add_run("Panduan Operasional  •  "), 8, False, MID_GRAY)
add_page_number(fp)


def title(text, subtitle=None):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run(text)
    if subtitle:
        p2 = doc.add_paragraph(style="Subtitle")
        p2.add_run(subtitle)


def h1(text):
    return doc.add_heading(text, level=1)


def h2(text):
    return doc.add_heading(text, level=2)


def h3(text):
    return doc.add_heading(text, level=3)


def para(text="", bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        font(p.add_run(bold_prefix), bold=True)
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def bullet(text):
    return doc.add_paragraph(text, style="List Bullet")


def step(text):
    p = doc.add_paragraph(text, style="List Number")
    p.paragraph_format.keep_together = True
    return p


def callout(label, text, color=BLUE, fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.55)
    row_pr = table.rows[0]._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    row_pr.append(cant_split)
    cell = table.cell(0, 0)
    cell.width = Inches(6.55)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, 140, 160, 140, 160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    font(p.add_run(label + "  "), bold=True, color=color)
    p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    t.style = "Table Grid"
    hdr = t.rows[0]
    set_repeat_table_header(hdr)
    for idx, text in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        font(p.add_run(text), 9.5, True, WHITE)
    for row_data in rows:
        cells = t.add_row().cells
        for idx, text in enumerate(row_data):
            cell = cells[idx]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if len(t.rows) % 2 == 1:
                set_cell_shading(cell, "F8FAFC")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            font(p.add_run(str(text)), 9.2)
    if widths:
        for row in t.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def screenshot(filename, caption, width=6.55):
    path = SHOTS / filename
    if not path.exists():
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    shape = run.add_picture(str(path), width=Inches(width))
    doc_pr = shape._inline.docPr
    doc_pr.set("descr", caption)
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(8)
    cp.paragraph_format.keep_together = True
    font(cp.add_run(caption), 8.5, False, MID_GRAY, italic=True)


def page_break():
    doc.add_page_break()


# Cover
doc.add_paragraph().paragraph_format.space_after = Pt(78)
kicker = doc.add_paragraph()
font(kicker.add_run("ADMINISTRATOR HANDBOOK"), 10, True, GOLD)
title("Panduan Penggunaan Admin CMS", "Mahaga Widya Cita • Artikel, Jurnal, Terjemahan, Media, Pengguna, dan Konten Website")
doc.add_paragraph().paragraph_format.space_after = Pt(35)
callout("Tujuan", "Panduan langkah demi langkah untuk mengelola website corporate secara aman, konsisten, dan bilingual.", NAVY, "EEF4F8")
meta = doc.add_paragraph()
meta.paragraph_format.space_before = Pt(30)
font(meta.add_run("Versi dokumentasi\n"), 9, True, MID_GRAY)
font(meta.add_run(date.today().strftime("%d %B %Y") + "\n"), 11, True, NAVY)
font(meta.add_run("Lingkungan: Production  |  Bahasa: Indonesia / English"), 9, False, MID_GRAY)
page_break()

h1("Daftar Isi")
for item in [
    "Mulai, login, dan mengenali dashboard",
    "Peran pengguna dan batas kewenangan",
    "Memilih Bahasa Indonesia atau Inggris",
    "Membuat dan menerbitkan artikel",
    "Menambahkan jurnal dan impor OJS",
    "Mengelola media dan upload gambar/PDF",
    "Menambahkan pengguna",
    "Menambahkan anggota tim",
    "Kategori, layanan, dan Policy Review",
    "Mengelola halaman corporate dan navigasi",
    "Pesan masuk, subscriber, dan privasi",
    "Checklist publikasi dan troubleshooting",
]:
    bullet(item)
callout("Prinsip utama", "Simpan konten Indonesia terlebih dahulu, periksa badge locale, baru jalankan terjemahan. Persetujuan terjemahan tidak otomatis memublikasikan konten.", GOLD, "FFF8E8")

h1("1. Mulai, Login, dan Dashboard")
h2("1.1 Login")
step("Buka https://www.mahagawidyacita.com/admin.")
step("Masukkan email dan password akun CMS.")
step("Setelah masuk, pastikan nama/avatar akun muncul di kanan atas.")
step("Gunakan Keluar / Log out setelah selesai, terutama pada perangkat bersama.")
callout("Keamanan", "Jangan membagikan password melalui dokumen, chat publik, atau screenshot. Gunakan akun pribadi dan prinsip least privilege.", RED, "FCECEC")

h2("1.2 Dashboard dan aksi cepat")
screenshot("01-dashboard-id.png", "Gambar 1. Dashboard Bahasa Indonesia dengan menu utama, locale, dan aksi cepat.")
para("Dashboard menyediakan akses cepat ke Artikel Baru, Jurnal Baru, Upload Media, Pesan Masuk, dan Impor Jurnal melalui link OJS. Ringkasan konten membantu admin melihat pekerjaan yang perlu ditindaklanjuti.")

h1("2. Peran Pengguna dan Kewenangan")
table(
    ["Peran", "Kewenangan utama", "Batas penting"],
    [
        ["Super Admin", "Seluruh konten, user, global, konfigurasi modul, publish", "Gunakan sangat terbatas"],
        ["Admin", "Kelola user, konten, review, publish, audience, media", "Tidak selalu mengubah konfigurasi khusus Super Admin"],
        ["Editor", "Membuat dan mengubah konten sendiri, media, kirim review", "Tidak dapat publish/schedule"],
        ["Reviewer", "Review, setujui, atau minta revisi", "Tidak membuat konten/publish"],
        ["Member (legacy)", "Read-only", "Jangan dipakai untuk akun baru"],
    ],
    widths=[1.05, 3.45, 2.05],
)
callout("Rekomendasi", "Pilih role paling rendah yang masih memungkinkan pekerjaan pengguna. Naikkan role hanya jika memang diperlukan.", GREEN, "ECF7F1")

h1("3. Bahasa Indonesia dan Inggris")
h2("3.1 Cara mengetahui locale aktif")
step("Periksa tombol Locale/Lokal di kanan atas. Nilainya harus id atau en.")
step("Periksa URL: ?locale=id untuk Indonesia dan ?locale=en untuk Inggris.")
step("Periksa badge workflow pada dokumen: ID / Mengedit Indonesia atau EN / Editing English.")
screenshot("02-dashboard-en.png", "Gambar 2. Dashboard dan navigasi admin dalam Bahasa Inggris ketika locale EN aktif.")

h2("3.2 Workflow terjemahan Indonesia ke Inggris")
step("Buat atau perbarui konten pada locale ID.")
step("Klik Simpan. Jangan menerjemahkan saat masih ada perubahan yang belum disimpan.")
step("Klik Buat Bahasa Inggris. Jika sumber berubah, gunakan Perbarui Bahasa Inggris.")
step("Tunggu status Draf AI siap direview.")
screenshot("03-category-id-translation.png", "Gambar 3. Status workflow terjemahan pada versi Indonesia.")
step("Pindah ke locale EN, lalu buka Tinjau draf AI.")
step("Bandingkan sumber Indonesia dengan kandidat Inggris: istilah korporat, nama orang, angka, URL, DOI/ISSN, dan rich text.")
step("Admin/Super Admin klik Setujui Bahasa Inggris jika hasil sudah benar.")
screenshot("04-category-en-review.png", "Gambar 4. Versi Inggris siap direview; approval tidak sama dengan publishing.")
callout("Batas saat ini", "Workflow otomatis yang tersedia adalah Indonesia ke Inggris. Pembuatan otomatis Inggris ke Indonesia belum tersedia.", GOLD, "FFF8E8")

h1("4. Membuat Artikel Baru")
h2("4.1 Mulai artikel")
step("Pastikan locale ID aktif.")
step("Pilih Dashboard > Artikel Baru, atau Manajemen Konten > Artikel > Buat Baru.")
screenshot("05-articles-list.png", "Gambar 5. Daftar Artikel dan tombol Buat Baru.")

h2("4.2 Isi konten penulisan")
step("Isi Judul Artikel (wajib).")
step("Isi Isi Artikel menggunakan rich-text editor. Gunakan tombol media editor untuk gambar inline.")
step("Isi Ringkasan Artikel maksimal 320 karakter untuk card, pencarian, dan metadata.")
step("Periksa slug agar singkat, unik, dan tidak berubah setelah dipublikasikan tanpa rencana redirect.")
screenshot("06-article-content.png", "Gambar 6. Tab Konten Penulisan: judul, rich text, ringkasan, SEO, slug, author, kategori, dan status.")

h2("4.3 Tambahkan gambar")
step("Buka tab Media.")
step("Pada Gambar Utama atau Meta Image, pilih Buat Baru untuk upload atau Pilih dari yang sudah ada untuk Media Library.")
step("Isi alt text yang menjelaskan isi gambar dan tambahkan caption/kredit bila diperlukan.")
step("Gunakan URL gambar lama hanya sebagai fallback konten legacy.")
screenshot("07-article-media.png", "Gambar 7. Tab Media artikel mendukung upload langsung dan pemilihan Media Library.")

h2("4.4 Review dan publikasi")
step("Simpan sebagai Draft.")
step("Buat dan review versi EN mengikuti Bagian 3.")
step("Editor klik Kirim Review.")
step("Reviewer memilih Setujui atau Minta Revisi dan memberi catatan yang spesifik.")
step("Admin memilih Publish atau menjadwalkan Published At.")
step("Buka Preview ID dan EN, lalu cek desktop serta mobile.")
table(
    ["Status", "Makna", "Tindakan berikutnya"],
    [
        ["Draft", "Belum siap direview", "Lengkapi isi dan media"],
        ["In Review", "Menunggu reviewer", "Review substansi dan bahasa"],
        ["Revision Requested", "Perlu perbaikan", "Editor revisi lalu kirim ulang"],
        ["Approved", "Lolos review", "Admin publish/schedule"],
        ["Published", "Tampil publik", "Monitor dan preview kedua locale"],
        ["Archived", "Dinonaktifkan", "Aktifkan ulang bila diperlukan"],
    ],
    widths=[1.35, 2.4, 2.8],
)

page_break()
h1("5. Menambahkan Jurnal")
h2("5.1 Membuat jurnal manual")
step("Pilih Dashboard > Jurnal Baru atau Manajemen Konten > Jurnal > Buat Baru.")
step("Isi Judul, Abstrak, dan minimal satu Kata Kunci.")
step("Isi ringkasan/isi tambahan bila tersedia.")
screenshot("08-journal-content.png", "Gambar 8. Tab Konten Jurnal untuk judul, abstrak, dan kata kunci.")

h2("5.2 Media dan metadata publikasi")
step("Buka tab Media & Publikasi.")
step("Upload Cover bila tersedia.")
step("Upload Dokumen Jurnal dalam format PDF - field ini wajib.")
step("Pilih Bahasa Publikasi: Indonesia, English, atau Bilingual.")
step("Isi Tahun Terbit, Volume, Issue, Halaman, DOI, ISSN, dan Link OJS bila ada.")
screenshot("09-journal-publication.png", "Gambar 9. Media dan metadata jurnal, termasuk upload PDF wajib.")

h2("5.3 Penulis")
step("Buka tab Penulis lalu tambah minimal satu penulis.")
step("Isi nama, afiliasi, email publik, dan URL profil jika tersedia.")
step("Atur urutan penulis sesuai kontribusi/publikasi asli.")
screenshot("10-journal-authors.png", "Gambar 10. Tab Penulis jurnal.")

h2("5.4 Impor dari OJS")
step("Dari Dashboard klik Impor Jurnal (Link OJS).")
step("Tempel URL artikel OJS, lalu jalankan impor.")
step("Verifikasi ulang judul, penulis, abstrak, DOI, tahun, volume, issue, dan halaman.")
step("Pastikan PDF, cover, kategori, locale, dan status sudah benar sebelum publish.")
screenshot("26-ojs-import.png", "Gambar 11. Dialog impor jurnal melalui link OJS.")
callout("Wajib diverifikasi", "Impor OJS mempercepat input, tetapi metadata hasil ekstraksi tetap harus dibandingkan dengan halaman dan PDF jurnal asli.", GOLD, "FFF8E8")

page_break()
h1("6. Media: Gambar dan PDF")
step("Pilih Media > Buat Baru atau Dashboard > Upload Media.")
step("Pilih/drag file gambar atau PDF.")
step("Untuk gambar, isi Alt Text wajib pada locale yang benar.")
step("Tambahkan caption dan kredit/sumber bila diperlukan.")
step("Simpan, lalu pilih media tersebut dari artikel, jurnal, tim, atau global.")
screenshot("11-media-upload.png", "Gambar 12. Form upload Media dengan area drag-and-drop dan metadata aksesibilitas.")
bullet("Gunakan nama file deskriptif dan ukuran yang wajar.")
bullet("Jangan menghapus media yang masih direferensikan konten.")
bullet("Sistem menerima image/* dan application/pdf serta membuat beberapa ukuran gambar.")

h1("7. Menambahkan Pengguna")
step("Pilih Collections/Koleksi > Users > Buat Baru.")
step("Isi Email, Password, Confirm Password, dan Role.")
step("Pilih Editor untuk penulis konten, Reviewer untuk pemeriksa, Admin untuk pengelola penuh, atau Super Admin hanya untuk pemilik sistem.")
step("Simpan lalu sampaikan URL login secara aman. Minta pengguna mengganti password.")
screenshot("13-user-create.png", "Gambar 13. Form membuat user baru. Password tidak pernah dicantumkan dalam dokumentasi.")
screenshot("redacted-12-users-list.png", "Gambar 14. Daftar user dengan data pribadi disensor.")
callout("Akses", "Hanya Admin/Super Admin yang seharusnya mengelola seluruh akun. Jangan gunakan role Member legacy untuk akun baru.", RED, "FCECEC")

page_break()
h1("8. Menambahkan Anggota Tim")
step("Pilih Manajemen Konten > Tim Ahli > Buat Baru.")
step("Pada Data Profil, upload foto, isi Nama Lengkap, Inisial (maks. 3), dan Bio.")
step("Pilih kategori Management atau Expert.")
step("Untuk Management isi Jabatan; untuk Expert isi Instansi Asal.")
step("Isi Bidang Keahlian, gradient, dan Urutan. Nomor lebih kecil tampil lebih atas.")
step("Simpan versi ID, buat EN, review jabatan/bio/keahlian, lalu approve.")
screenshot("14-team-list.png", "Gambar 15. Daftar anggota tim.")
screenshot("15-team-create.png", "Gambar 16. Form profil anggota tim.")
callout("Pemeriksaan terjemahan", "Nama orang dan nama institusi biasanya tidak diterjemahkan. Fokuskan review pada jabatan, bio, dan bidang keahlian.", GREEN, "ECF7F1")

h1("9. Kategori, Layanan, dan Policy Review")
h2("9.1 Kategori")
step("Buat Name pada locale ID dan slug unik.")
step("Simpan, buat versi EN, review, dan approve.")
step("Jangan mengubah slug yang sudah dipakai tanpa rencana migrasi URL.")
screenshot("16-categories.png", "Gambar 17. Daftar kategori konten.")

h2("9.2 Layanan")
step("Isi Judul, slug, tagline, deskripsi, tema/gradient, features, benefits, dan target audience.")
step("Simpan ID, buat dan review EN.")
step("Uji halaman publik setelah perubahan slug atau tema.")
screenshot("17-services.png", "Gambar 18. Daftar layanan corporate.")

h2("9.3 Policy Review")
para("Policy Review mengikuti pola editorial yang mirip artikel/jurnal dan membutuhkan dokumen. Modul ini dapat disembunyikan melalui Pengaturan Fitur & Modul oleh Super Admin.")
screenshot("25-feature-settings.png", "Gambar 19. Pengaturan fitur untuk mengaktifkan/menonaktifkan modul Policy Review.")

page_break()
h1("10. Mengelola Halaman Corporate")
h2("10.1 Beranda")
bullet("Atur visibility section: Hero, Statistik, Mitra, Layanan, Artikel, Tim, dan CTA.")
bullet("Kelola hero, statistik, mitra/logo, layanan unggulan, artikel, tim, WhatsApp CTA, dan benefit points.")
bullet("Preview kedua locale karena perubahan global berdampak luas.")
screenshot("18-home-settings.png", "Gambar 20. Pengaturan Beranda dan visibility section.")

h2("10.2 Tentang Kami dan Kontak")
bullet("Tentang Kami: hero, statistik, profil, visi, misi, nilai inti, dan pesan CEO.")
bullet("Kontak: telepon/WhatsApp, email, alamat, jam operasional, lokasi, CTA, dan subjek form.")
screenshot("19-about-settings.png", "Gambar 21. Pengaturan halaman Tentang Kami.")
screenshot("20-contact-settings.png", "Gambar 22. Pengaturan Informasi Kontak.")

h2("10.3 Navbar dan Footer")
step("Navbar: isi label dan href; gunakan children untuk dropdown.")
step("Pastikan path internal diawali /id atau /en sesuai desain routing.")
step("Footer: kelola deskripsi perusahaan, social URL, link perusahaan/layanan, dan copyright.")
step("Uji desktop/mobile serta locale ID/EN setelah menyimpan.")
screenshot("21-navbar-settings.png", "Gambar 23. Pengaturan Navbar dan menu atas.")
screenshot("22-footer-settings.png", "Gambar 24. Pengaturan Footer.")

page_break()
h1("11. Pesan Masuk, Subscribers, dan Privasi")
h2("11.1 Pesan Masuk")
bullet("Baca nama, email, telepon, institusi, subjek, dan pesan dari form publik.")
bullet("Jangan mengubah isi asli kecuali ada kebutuhan operasional yang jelas.")
bullet("Batasi akses karena informasi dapat mengandung data pribadi.")
screenshot("redacted-23-contact-submissions.png", "Gambar 25. Daftar pesan masuk; data pribadi disensor.")

h2("11.2 Subscribers")
bullet("Email subscriber berasal dari newsletter publik dan harus unik.")
bullet("Kelola export/delete sesuai kebijakan privasi dan permintaan unsubscribe.")
screenshot("redacted-24-subscribers.png", "Gambar 26. Daftar subscriber; alamat email disensor.", width=5.45)

h1("12. Checklist Sebelum Publish")
for item in [
    "Badge locale sesuai: ID untuk sumber, EN untuk versi Inggris.",
    "Judul, slug, kategori, author/penanggung jawab, dan status benar.",
    "Gambar utama/cover tampil dan memiliki alt text.",
    "Jurnal memiliki PDF dan minimal satu penulis.",
    "Link, CTA WhatsApp, DOI/ISSN, dan URL eksternal dapat dibuka.",
    "Terjemahan sudah direview dan di-approve oleh role berwenang.",
    "Konten sudah melalui workflow review; approval translation bukan publish.",
    "Preview ID/EN diperiksa pada desktop dan mobile.",
    "Published At dan timezone sudah benar bila dijadwalkan.",
]:
    bullet("☐ " + item)

h2("Troubleshooting singkat")
table(
    ["Masalah", "Pemeriksaan", "Tindakan"],
    [
        ["Tidak tahu locale aktif", "Badge dan ?locale=", "Pilih id/en sebelum mengedit"],
        ["Tombol translate tidak aktif", "Perubahan belum disimpan", "Simpan lalu coba kembali"],
        ["Terjemahan gagal", "Status/error workflow", "Coba Lagi atau hubungi admin"],
        ["Tidak dapat publish", "Role dan status review", "Minta Admin/Super Admin"],
        ["Gambar tidak tampil", "Media, alt, relasi", "Upload/pilih ulang lalu Save"],
        ["Jurnal gagal disimpan", "PDF, penulis, tahun", "Lengkapi seluruh field wajib"],
        ["URL berubah/rusak", "Slug diedit", "Kembalikan slug atau siapkan redirect"],
    ],
    widths=[1.7, 2.3, 2.55],
)

doc.core_properties.title = "Panduan Penggunaan Admin CMS Mahaga Widya Cita"
doc.core_properties.subject = "Manual operasional admin bilingual"
doc.core_properties.author = "Mahaga Widya Cita"
doc.core_properties.keywords = "CMS, admin, artikel, jurnal, terjemahan, media"
doc.save(DOCX_PATH)
print(DOCX_PATH)
